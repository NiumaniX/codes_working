import pandas as pd

from 周报自动化_functions import tool
import os
from docxtpl import DocxTemplate
import warnings
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
warnings.filterwarnings("ignore")
from docx.shared import Inches
####################################
######################################################################################
                                        #''' 一、 产品管理总体情况｜ 二、 新股投资策略情况'''
invest_of_return_info = tool.read_excel(os.getcwd() + '/区间打新-产品汇总情况.xlsx')
invest_of_return_info = invest_of_return_info.drop(invest_of_return_info.columns[[0]], axis = 1).drop([0], axis = 0)
#print(invest_of_return_info.info())
                                        #''' 全局变量1 --- 时间'''
year, input_date = '2022','0916'  #修改
#input_date = input('请输入周报的产生时间(e.g. 0902)')
month_str, day_str = eval('str(int(input_date[:2]))'), eval('str(int(input_date[2:4]))')

                                        #''' 全局变量2 --- 基金名称'''
funds_list =  ['XXX单一资产管理计划',
               ]
                                        #''' 全局变量3 --- 投资经理'''
manager_list = {'XXX单一资产管理计划':'XXX',
                }
                                        #''' 全局变量4 --- 净值截止日'''
net_value_closingday = tool.format_datetime(year + input_date, '%Y/%m/%d', 1, True)
print('strf_time:', net_value_closingday)
                                        #''' 全局变量5 --- 周报起始日与截止日'''
start_date_of_weekly_report = tool.format_datetime(year + input_date, '%m月%d日', 5, True)
end_date_of_weekly_report = tool.format_datetime(year + input_date, '%m月%d日', 0, True)
print('start_date_of_weekly_report:', start_date_of_weekly_report)
print('end_date_of_weekly_report:', end_date_of_weekly_report)




                                       #''' 一、 产品管理总体情况'''


                                       #''' 二、 新股投资策略情况'''
#print(invest_of_return_info.head())
########

for fund_name in funds_list:
                                        #''' 全局变量6 --- ipo_num_of_funds'''
    ipo_num_of_funds = invest_of_return_info[['产品名称','科创板入围个数', '创业板入围个数', '主板入围个数']][invest_of_return_info['产品名称'] == fund_name].sum(axis = 1).values[0]
    # print(ipo_num_of_funds)
                                        #''' 全局变量7 --- lucky_money'''
    lucky_money = invest_of_return_info[['产品名称','区间获配金额']][invest_of_return_info['产品名称'] == fund_name].values[0][1]
    lucky_money = round(lucky_money/10000, 2)
    # print('lucky_money',lucky_money)
                                        #''' 全局变量8 --- lucky_returns'''
    lucky_returns = invest_of_return_info[['产品名称','区间卖出收益']][invest_of_return_info['产品名称'] == fund_name].values[0][1]
    lucky_returns = round(lucky_returns/10000, 2)
    # print('lucky_returns',lucky_returns)
    #                                     #''' 全局变量9 --- accumulated_chosen_IPO_num'''
    #accumulated_chosen_IPO_num = invest_of_return_info[['产品名称','累计获配金额']][invest_of_return_info['产品名称'] == fund_name].values[0][1]
    accumulated_chosen_IPO_num = '【暂无】'
                                        #''' 全局变量10 --- accumulated_lucky_money'''
    accumulated_lucky_money = invest_of_return_info[['产品名称','累计获配金额']][invest_of_return_info['产品名称'] == fund_name].values[0][1]
    accumulated_lucky_money = round(accumulated_lucky_money/10000, 2)
    #print('accumulated_lucky_money',accumulated_lucky_money)
                                        #''' 全局变量11 --- accumulated_lucky_returns'''
    accumulated_lucky_returns = invest_of_return_info[['产品名称','累计卖出收益']][invest_of_return_info['产品名称'] == fund_name].values[0][1]
    accumulated_lucky_returns = round(accumulated_lucky_returns/10000, 2)
    # print('accumulated_lucky_returns',accumulated_lucky_returns)
    # print('#'*30)

    # context = { '1.month_str': month_str, '1.day_str': day_str,
    #             '2.fund_name': fund_name,
    #             '3.fund_manager': manager_list[fund_name],
    #             '4.净值截止日': net_value_closingday,
    #             '5.周报的起始日': start_date_of_weekly_report, '5.周报的截止日': end_date_of_weekly_report,
    #             '6.ipo_num_of_funds':ipo_num_of_funds,
    #             '7.lucky_money':lucky_money,
    #             '8.lucky_returns':lucky_returns,
    #             '9.accumulated_chosen_IPO_num': accumulated_chosen_IPO_num,
    #             '10.accumulated_lucky_money': accumulated_lucky_money,
    #             '11.accumulated_lucky_returns': accumulated_lucky_returns
    #         }
    file_path = os.getcwd() +  r'/周报自动化-set'
    try:
      os.mkdir(file_path)  #创建一级目录
    except:
      pass

    doc = DocxTemplate("周报自动化-模版.docx")
    context = { 'month_str': month_str, 'day_str': day_str,
            'fund_name': fund_name,
            'fund_manager': manager_list[fund_name],
            '净值截止日': net_value_closingday,
            '周报的起始日': start_date_of_weekly_report, '周报的截止日': end_date_of_weekly_report,
            'chosen_IPO_num':ipo_num_of_funds,
            'lucky_money':lucky_money,
            'lucky_returns':lucky_returns,
            #'accumulated_chosen_IPO_num': accumulated_chosen_IPO_num,
            'accumulated_lucky_money': accumulated_lucky_money,
            'accumulated_lucky_returns': accumulated_lucky_returns}
    doc.render(context)
    df = pd.read_excel(os.getcwd()+'/区间打新-产品明细情况.xlsx')
    df = df.drop(df.columns[[0]], axis = 1)
    columns = ['市场名称',	'股票代码',	'股票名称',	'发行价格',	'获配数量'	,
               '获配金额',	'卖出数量',	'卖出收益']
    df = df[df["产品名称"] == fund_name][columns]
    print(df)
    # add the table.
    t = doc.add_table(df.shape[0]+1, df.shape[1])
    t.style = 'Table Grid'
    # add the header rows.
    for j in range(df.shape[-1]):
        t.cell(0,j).text = df.columns[j]

    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i+1,j).text = str(df.values[i,j])
            t.cell(i+1,j).paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for cell in t.columns[0].cells:
        cell.width = Inches(0.8)
    doc.save(file_path + r"/{}-20220916.docx".format(fund_name))



